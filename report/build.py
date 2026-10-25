# -*- coding: utf-8 -*-
"""Dựng báo cáo .docx từ các tệp nội dung content1..content5."""

import importlib
import sys
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

FONT = 'Times New Roman'
MONO = 'Consolas'


# ---------------------------------------------------------------- helpers
def set_run_font(run, name=FONT, size=13, bold=None, italic=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), name)


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def add_field(paragraph, instr, placeholder=''):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr_el = OxmlElement('w:instrText')
    instr_el.set(qn('xml:space'), 'preserve')
    instr_el.text = instr
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    text_el = OxmlElement('w:t')
    text_el.text = placeholder
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_sep)
    run._r.append(text_el)
    run._r.append(fld_end)
    return run


def para(doc, text='', size=13, bold=False, italic=False, align=None,
         first_indent=None, left_indent=None, space_before=0, space_after=6,
         line=1.3, style=None, keep=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if left_indent is not None:
        pf.left_indent = left_indent
    if keep:
        pf.keep_with_next = True
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def doc_paragraph_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.3
    set_run_font(p.add_run(text), size=16, bold=True)
    return p


# ---------------------------------------------------------------- styles
def strip_theme_fonts(rPr):
    """Bỏ thuộc tính font theo chủ đề để Word dùng đúng Times New Roman."""
    if rPr is None:
        return
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rFonts.get(qn(attr)) is not None:
                del rFonts.attrib[qn(attr)]
    color = rPr.find(qn('w:color'))
    if color is not None:
        for attr in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
            if color.get(qn(attr)) is not None:
                del color.attrib[qn(attr)]


def setup_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(13)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), FONT)
    strip_theme_fonts(rpr)
    pf = normal.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(6)

    specs = {
        'Heading 1': (16, True, 12, 8),
        'Heading 2': (14, True, 10, 6),
        'Heading 3': (13, True, 8, 4),
    }
    for name, (size, bold, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        r = st.element.get_or_add_rPr()
        rf = r.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            r.append(rf)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(attr), FONT)
        strip_theme_fonts(r)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def setup_page(doc):
    s = doc.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(3.0)
    s.right_margin = Cm(2.0)
    s.different_first_page_header_footer = True

    footer = s.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, 'PAGE', '1')
    for run in p.runs:
        set_run_font(run, size=12)
    return s


# ---------------------------------------------------------------- blocks
class Builder:
    def __init__(self, doc):
        self.doc = doc
        self.num = 0
        self.col_widths = None

    def reset_num(self):
        self.num = 0

    def h1(self, text):
        p = self.doc.add_paragraph(text, style='Heading 1')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_num()

    def h1c(self, text):
        p = self.doc.add_paragraph(text, style='Heading 1')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_num()

    def h2(self, text):
        self.doc.add_paragraph(text, style='Heading 2')
        self.reset_num()

    def h3(self, text):
        self.doc.add_paragraph(text, style='Heading 3')
        self.reset_num()

    def p(self, text):
        para(self.doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_indent=Cm(1.0), space_after=6)

    def label(self, text):
        p = para(self.doc, text, bold=True, italic=False, space_before=8,
                 space_after=3, keep=True)
        self.reset_num()
        return p

    def mt(self, code, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.4)
        pf.first_line_indent = Cm(-1.4)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(p.add_run(code + '. '), bold=True)
        set_run_font(p.add_run(text))
        self.reset_num()

    def bullet(self, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.4)
        pf.first_line_indent = Cm(-0.6)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(p.add_run('– '))
        set_run_font(p.add_run(text))
        self.reset_num()

    def numbered(self, text):
        self.num += 1
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.4)
        pf.first_line_indent = Cm(-0.7)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(p.add_run('%d. ' % self.num))
        set_run_font(p.add_run(text))

    def table(self, caption, headers, rows, widths):
        total = sum(widths)
        if total <= 0:
            total = 1.0
        factor = 15.8 / total
        widths = [w * factor for w in widths]
        fs = 12
        if len(headers) >= 7:
            fs = 9.5
        elif len(headers) == 6:
            fs = 10.5
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.line_spacing = 1.15
        cap.paragraph_format.keep_with_next = True
        num, _, rest = caption.partition('.')
        set_run_font(cap.add_run(num + '.'), size=12, bold=True)
        set_run_font(cap.add_run(rest), size=12)

        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_cell_margins(t)
        hdr = t.rows[0]
        repeat_header(hdr)
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(h), size=fs, bold=True)
            shade_cell(cell, 'EDEDED')

        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cell = cells[i]
                cell.text = ''
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                if len(str(val)) <= 6:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run_font(p.add_run(str(val)), size=fs)

        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)

        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(4)
        after.paragraph_format.line_spacing = 1.0
        self.reset_num()

    def figure(self, num, caption, hint):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        shade_cell(cell, 'F4F4F4')
        cell.width = Cm(15.5)
        cell.text = ''
        p0 = cell.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(60)
        p0.paragraph_format.space_after = Pt(6)
        set_run_font(p0.add_run('« CHÈN ' + num.upper() + ' TẠI ĐÂY »'), size=12, bold=True,
                     color=RGBColor(0x88, 0x88, 0x88))
        p1 = cell.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(60)
        p1.paragraph_format.line_spacing = 1.1
        set_run_font(p1.add_run('Gợi ý nội dung: ' + hint), size=10.5, italic=True,
                     color=RGBColor(0x77, 0x77, 0x77))

        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.line_spacing = 1.15
        set_run_font(cap.add_run(num + '. '), size=12, bold=True)
        set_run_font(cap.add_run(caption), size=12, italic=True)
        self.reset_num()

    def code(self, caption, text):
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(2)
        cap.paragraph_format.line_spacing = 1.1
        cap.paragraph_format.keep_with_next = True
        set_run_font(cap.add_run(caption), size=11, italic=True)

        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.4)
        pf.right_indent = Cm(0.2)
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.0
        shade_paragraph(p, 'F5F5F5')
        lines = text.split('\n')
        for i, line in enumerate(lines):
            run = p.add_run(line)
            set_run_font(run, name=MONO, size=9)
            if i < len(lines) - 1:
                run.add_break()

        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(6)
        after.paragraph_format.line_spacing = 1.0
        self.reset_num()

    def pre(self, text):
        self.code('Cấu trúc thư mục (rút gọn):', text)

    def refs(self, items):
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.9)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.25
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_run_font(p.add_run('[%d] ' % i))
            set_run_font(p.add_run(item), size=12)

    def toc(self, title):
        p = doc_paragraph_title(self.doc, title)

        holder = self.doc.add_paragraph()
        holder.paragraph_format.space_before = Pt(6)
        run = add_field(holder, 'TOC \\o "1-3" \\h \\z \\u',
                        'Mục lục sẽ hiện ở đây sau khi cập nhật trường (F9).')
        set_run_font(run, size=12)

        note = self.doc.add_paragraph()
        note.paragraph_format.space_before = Pt(10)
        set_run_font(note.add_run('Ghi chú khi hoàn thiện: bấm chuột phải vào vùng mục lục → '
                                  'Update Field (hoặc F9) để Word tính lại số trang sau khi '
                                  'đã chèn toàn bộ hình và bảng.'),
                     size=10.5, italic=True, color=RGBColor(0x77, 0x77, 0x77))
        self.reset_num()


# ---------------------------------------------------------------- cover
def build_cover(doc, cover):
    para(doc, cover['org1'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=18, space_after=2)
    para(doc, cover['org2'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=24)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = line._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F3864')
    pbdr.append(bottom)
    pPr.append(pbdr)
    line.paragraph_format.space_after = Pt(26)

    para(doc, cover['kind'], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=30)

    para(doc, cover['label'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=8)
    para(doc, cover['title'], size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=30, line=1.35)

    para(doc, cover['field'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=10, space_after=40)

    para(doc, 'Tác giả:', size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=6)
    for name, role in cover['authors']:
        para(doc, name, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        para(doc, '(%s)' % role, size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=8)

    for _ in range(3):
        para(doc, '', space_after=10)
    para(doc, cover['place'], size=13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=40)


# ---------------------------------------------------------------- main
def main():
    mods = {i: importlib.import_module('content%d' % i) for i in (1, 2, 3, 4, 5) if i != 3}
    # content3a / content3b: phần module, chèn sau content2
    c3a = importlib.import_module('content3a')
    c3b = importlib.import_module('content3b')

    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    build_cover(doc, mods[1].COVER)

    blocks = []
    blocks += mods[1].BLOCKS
    blocks += mods[2].BLOCKS
    blocks += c3a.BLOCKS
    blocks += c3b.BLOCKS
    blocks += mods[4].BLOCKS
    blocks += mods[5].BLOCKS

    b = Builder(doc)
    for block in blocks:
        kind = block[0]
        if kind == 'h1':
            b.h1(block[1])
        elif kind == 'h1c':
            b.h1c(block[1])
        elif kind == 'h2':
            b.h2(block[1])
        elif kind == 'h3':
            b.h3(block[1])
        elif kind == 'p':
            b.p(block[1])
        elif kind == 'label':
            b.label(block[1])
        elif kind == 'mt':
            b.mt(block[1], block[2])
        elif kind == 'bullet':
            b.bullet(block[1])
        elif kind == 'numbered':
            b.numbered(block[1])
        elif kind == 'table':
            b.table(block[1], block[2], block[3], block[4])
        elif kind == 'figure':
            b.figure(block[1], block[2], block[3])
        elif kind == 'code':
            b.code(block[1], block[2])
        elif kind == 'pre':
            b.pre(block[1])
        elif kind == 'refs':
            b.refs(block[1])
        elif kind == 'toc':
            b.toc(block[1])
        elif kind == 'pagebreak':
            doc.add_page_break()
        else:
            raise ValueError('Block không hợp lệ: %r' % (kind,))

    out = os.path.join(os.path.dirname(HERE), 'BAO_CAO_ELEC_AI.docx')
    doc.save(out)
    print('saved:', out)
    print('blocks:', len(blocks))


if __name__ == '__main__':
    main()
