# -*- coding: utf-8 -*-
import re
import zipfile
import docx

P = r'D:\Tải Xuống\app\báo cáo\BAO_CAO_ELEC_AI.docx'
d = docx.Document(P)

# 1. token còn sót
txt = '\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            txt += '\n' + c.text
left = re.findall(r'\{(?:fig|tab):[a-z0-9]+\}', txt)
print('token còn sót:', left)

# 2. đánh số hình/bảng
figs = [int(m.group(1)) for p in d.paragraphs
        for m in [re.match(r'Hình (\d+)\.', p.text.strip())] if m]
tabs = [int(m.group(1)) for p in d.paragraphs
        for m in [re.match(r'Bảng (\d+)\.', p.text.strip())] if m]
print('hình:', len(figs), 'liên tục:', figs == list(range(1, len(figs) + 1)))
print('bảng:', len(tabs), 'liên tục:', tabs == list(range(1, len(tabs) + 1)))

# 3. tham chiếu chéo có trỏ tới số không tồn tại
refs_f = sorted({int(x) for x in re.findall(r'(?:hình|Hình) (\d+)', txt)})
refs_t = sorted({int(x) for x in re.findall(r'(?:bảng|Bảng) (\d+)', txt)})
print('tham chiếu hình vượt quá:', [n for n in refs_f if n > len(figs)])
print('tham chiếu bảng vượt quá:', [n for n in refs_t if n > len(tabs)])

# 4. bí mật
for pat in ('AIza', 'sk-', 'eyJ', 'serviceAccount', 'Bearer '):
    if pat in txt:
        print('CẢNH BÁO chuỗi:', pat)
print('số dài >=9 chữ số:', len(re.findall(r'\b\d{9,}\b', txt)))

# 5. định dạng
sec = d.sections[0]
p0 = [p for p in d.paragraphs if len(p.text) > 300][0]
print('trang:', round(sec.page_width.cm, 1), 'x', round(sec.page_height.cm, 1),
      '| lề:', round(sec.top_margin.cm, 1), round(sec.bottom_margin.cm, 1),
      round(sec.left_margin.cm, 1), round(sec.right_margin.cm, 1))
print('thân bài:', p0.runs[0].font.name, p0.runs[0].font.size.pt,
      '| giãn dòng:', p0.paragraph_format.line_spacing,
      '| canh đều:', p0.alignment)
z = zipfile.ZipFile(P)
styles = z.read('word/styles.xml').decode('utf-8')
print('theme font còn trong Heading1:', 'asciiTheme' in re.search(
    r'<w:style [^>]*w:styleId="Heading1".*?</w:style>', styles, re.S).group(0))
print('có trường TOC:', 'TOC ' in z.read('word/document.xml').decode('utf-8'))
print('tệp nặng (KB):', round(len(open(P, 'rb').read()) / 1024, 1))
print('số đoạn:', len(d.paragraphs), '| số bảng:', len(d.tables))
