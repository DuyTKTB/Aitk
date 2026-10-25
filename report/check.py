# -*- coding: utf-8 -*-
import re
import zipfile
import docx

z = zipfile.ZipFile(r'D:\Tải Xuống\app\BAO_CAO_ELEC_AI.docx')
styles = z.read('word/styles.xml').decode('utf-8')
for sid in ('Normal', 'Heading1', 'Heading2', 'Heading3'):
    m = re.search(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % sid, styles, re.S)
    body = m.group(0) if m else 'MISSING'
    has_theme = 'asciiTheme' in body or 'themeColor' in body
    fonts = re.findall(r'w:ascii="([^"]+)"', body)
    sizes = re.findall(r'<w:sz w:val="(\d+)"', body)
    print(sid, '| theme left:', has_theme, '| font:', fonts, '| sz(half-pt):', sizes)

d = docx.Document(r'D:\Tải Xuống\app\BAO_CAO_ELEC_AI.docx')
sec = d.sections[0]
print('page cm:', round(sec.page_width.cm, 1), 'x', round(sec.page_height.cm, 1))
print('margins T/B/L/R:', round(sec.top_margin.cm, 1), round(sec.bottom_margin.cm, 1),
      round(sec.left_margin.cm, 1), round(sec.right_margin.cm, 1))
p = [x for x in d.paragraphs if x.style.name == 'Normal' and len(x.text) > 200][0]
print('body: align', p.alignment, '| line', p.paragraph_format.line_spacing,
      '| indent', round(p.paragraph_format.first_line_indent.cm, 2))
print('first body run font:', p.runs[0].font.name, p.runs[0].font.size.pt)
t = d.tables[1]
print('table1 cols:', len(t.columns), 'rows:', len(t.rows),
      '| width col0 cm:', round(t.rows[0].cells[0].width.cm, 2))
print('total tables:', len(d.tables))
