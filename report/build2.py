# -*- coding: utf-8 -*-
"""Dựng báo cáo .docx (bản mở rộng) — đánh số hình/bảng tự động theo khóa."""

import importlib
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

FONT = 'Times New Roman'
MONO = 'Consolas'
MODULES = ['c%02d' % i for i in range(1, 20)]


# ---------------------------------------------------------------- low level
def set_run_font(run, name=FONT, size=13, bold=None, italic=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), name)


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_margins(table, top=30, bottom=30, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def keep_row_together(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


_BM_ID = [1000]


def add_bookmark(paragraph, name):
    """Đánh dấu đoạn để danh mục hình/bảng lấy được số trang."""
    _BM_ID[0] += 1
    bid = str(_BM_ID[0])
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), bid)
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_field(paragraph, instr, placeholder=''):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar')
    b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText')
    i.set(qn('xml:space'), 'preserve')
    i.text = instr
    s = OxmlElement('w:fldChar')
    s.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = placeholder
    e = OxmlElement('w:fldChar')
    e.set(qn('w:fldCharType'), 'end')
    for el in (b, i, s, t, e):
        run._r.append(el)
    return run


def strip_theme_fonts(rPr):
    if rPr is None:
        return
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rFonts.get(qn(attr)) is not None:
                del rFonts.attrib[qn(attr)]
    color = rPr.find(qn('w:color'))
    if color is not None:
        for attr in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
            if color.get(qn(attr)) is not None:
                del color.attrib[qn(attr)]


def para(doc, text='', size=13, bold=False, italic=False, align=None,
         first_indent=None, left_indent=None, space_before=0, space_after=6,
         line=1.3, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if left_indent is not None:
        pf.left_indent = left_indent
    if keep:
        pf.keep_with_next = True
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def setup_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(13)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(attr), FONT)
    strip_theme_fonts(rpr)
    pf = normal.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(6)

    specs = {'Heading 1': (16, 12, 8), 'Heading 2': (14, 10, 6), 'Heading 3': (13, 8, 4)}
    for name, (size, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        r = st.element.get_or_add_rPr()
        rf = r.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            r.append(rf)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(attr), FONT)
        strip_theme_fonts(r)
        spf = st.paragraph_format
        spf.space_before = Pt(before)
        spf.space_after = Pt(after)
        spf.line_spacing = 1.3
        spf.keep_with_next = True
        spf.alignment = WD_ALIGN_PARAGRAPH.LEFT


def setup_page(doc):
    s = doc.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(3.0)
    s.right_margin = Cm(2.0)
    s.different_first_page_header_footer = True
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, 'PAGE', '1')
    for run in p.runs:
        set_run_font(run, size=12)


def doc_paragraph_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.3
    set_run_font(p.add_run(text), size=16, bold=True)
    return p


# ---------------------------------------------------------------- builder
class Builder:
    def __init__(self, doc, fig_no, tab_no, fig_titles, tab_titles):
        self.doc = doc
        self.fig_no = fig_no
        self.tab_no = tab_no
        self.fig_titles = fig_titles
        self.tab_titles = tab_titles
        self.num = 0

    def sub(self, text):
        if not text:
            return text
        for key, n in self.fig_no.items():
            text = text.replace('{fig:%s}' % key, str(n))
        for key, n in self.tab_no.items():
            text = text.replace('{tab:%s}' % key, str(n))
        return text

    def reset_num(self):
        self.num = 0

    def h1(self, text):
        p = self.doc.add_paragraph(self.sub(text), style='Heading 1')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_num()

    def h1c(self, text):
        self.h1(text)

    def h2(self, text):
        self.doc.add_paragraph(self.sub(text), style='Heading 2')
        self.reset_num()

    def h3(self, text):
        self.doc.add_paragraph(self.sub(text), style='Heading 3')
        self.reset_num()

    def p(self, text):
        para(self.doc, self.sub(text), align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_indent=Cm(1.0), space_after=6)

    def label(self, text):
        para(self.doc, self.sub(text), bold=True, space_before=8, space_after=3, keep=True)
        self.reset_num()

    def mt(self, code, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.4)
        pf.first_line_indent = Cm(-1.4)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(p.add_run(code + '. '), bold=True)
        set_run_font(p.add_run(self.sub(text)))
        self.reset_num()

    def bullet(self, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.4)
        pf.first_line_indent = Cm(-0.6)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(p.add_run('– '))
        set_run_font(p.add_run(self.sub(text)))
        self.reset_num()

    def numbered(self, text):
        self.num += 1
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.4)
        pf.first_line_indent = Cm(-0.7)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(p.add_run('%d. ' % self.num))
        set_run_font(p.add_run(self.sub(text)))

    def spacer(self):
        para(self.doc, '', space_after=10)
        self.reset_num()

    def right(self, text, bold=False):
        para(self.doc, self.sub(text), bold=bold, align=WD_ALIGN_PARAGRAPH.RIGHT,
             space_after=4)
        self.reset_num()

    def note(self, text):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(1.0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.25
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(p.add_run(self.sub(text)), size=12, italic=True,
                     color=RGBColor(0x44, 0x44, 0x44))
        self.reset_num()

    def table(self, key, caption, headers, rows, widths):
        n = self.tab_no[key]
        total = sum(widths) or 1.0
        widths = [w * 15.8 / total for w in widths]
        fs = 12
        if len(headers) >= 7:
            fs = 9.5
        elif len(headers) == 6:
            fs = 10.5

        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(8)
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.line_spacing = 1.15
        cap.paragraph_format.keep_with_next = True
        set_run_font(cap.add_run('Bảng %d. ' % n), size=12, bold=True)
        set_run_font(cap.add_run(self.sub(caption)), size=12)
        add_bookmark(cap, 'Bang%d' % n)

        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_cell_margins(t)
        hdr = t.rows[0]
        repeat_header(hdr)
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(self.sub(h)), size=fs, bold=True)
            shade_cell(cell, 'EDEDED')

        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cell = cells[i]
                cell.text = ''
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                txt = self.sub(str(val))
                if len(txt) <= 8:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run_font(p.add_run(txt), size=fs)

        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)

        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(4)
        after.paragraph_format.line_spacing = 1.0
        self.reset_num()

    def figure(self, key, caption, hint):
        n = self.fig_no[key]
        t = self.doc.add_table(rows=1, cols=1)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        shade_cell(cell, 'F5F5F5')
        cell.width = Cm(15.5)
        cell.text = ''
        p0 = cell.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(78)
        p0.paragraph_format.space_after = Pt(6)
        set_run_font(p0.add_run('« CHÈN HÌNH %d TẠI ĐÂY »' % n), size=12, bold=True,
                     color=RGBColor(0x88, 0x88, 0x88))
        p1 = cell.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(78)
        p1.paragraph_format.line_spacing = 1.1
        set_run_font(p1.add_run('Gợi ý: ' + self.sub(hint)), size=10.5, italic=True,
                     color=RGBColor(0x77, 0x77, 0x77))

        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.line_spacing = 1.15
        set_run_font(cap.add_run('Hình %d. ' % n), size=12, bold=True)
        set_run_font(cap.add_run(self.sub(caption)), size=12, italic=True)
        add_bookmark(cap, 'Hinh%d' % n)
        self.reset_num()

    def code(self, caption, text):
        cap = self.doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(2)
        cap.paragraph_format.line_spacing = 1.1
        cap.paragraph_format.keep_with_next = True
        set_run_font(cap.add_run(self.sub(caption)), size=11, italic=True)

        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.4)
        pf.right_indent = Cm(0.2)
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.0
        shade_paragraph(p, 'F5F5F5')
        lines = text.split('\n')
        for i, line in enumerate(lines):
            run = p.add_run(line)
            set_run_font(run, name=MONO, size=9)
            if i < len(lines) - 1:
                run.add_break()

        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(6)
        after.paragraph_format.line_spacing = 1.0
        self.reset_num()

    def pre(self, text):
        self.code('Cấu trúc thư mục (rút gọn):', text)

    def refs(self, items):
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.9)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.25
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_run_font(p.add_run('[%d] ' % i))
            set_run_font(p.add_run(item), size=12)

    def _index_list(self, title, prefix, titles, total, bm):
        self.h1(title)
        intro = self.doc.add_paragraph()
        intro.paragraph_format.space_after = Pt(8)
        intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_run_font(intro.add_run(
            'Danh mục dưới đây liệt kê toàn bộ %s trong báo cáo. Cột số trang là trường '
            'tự động: sau khi chèn đủ hình và bấm cập nhật (F9), Word sẽ điền số trang đúng.'
            % ('hình vẽ' if prefix == 'Hình' else 'bảng biểu')), size=12, italic=True)

        for n in range(1, total + 1):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(2)
            pf.line_spacing = 1.2
            pf.left_indent = Cm(1.2)
            pf.first_line_indent = Cm(-1.2)
            pf.tab_stops.add_tab_stop(Cm(15.6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            set_run_font(p.add_run('%s %d. ' % (prefix, n)), size=12, bold=True)
            set_run_font(p.add_run(titles.get(n, '')), size=12)
            p.add_run().add_tab()
            run = add_field(p, 'PAGEREF %s%d \\h' % (bm, n), '…')
            set_run_font(run, size=12)
        self.reset_num()

    def figlist(self, title):
        self._index_list(title, 'Hình', self.fig_titles, len(self.fig_no), 'Hinh')

    def tablist(self, title):
        self._index_list(title, 'Bảng', self.tab_titles, len(self.tab_no), 'Bang')

    def toc(self, title):
        doc_paragraph_title(self.doc, title)
        holder = self.doc.add_paragraph()
        holder.paragraph_format.space_before = Pt(6)
        run = add_field(holder, 'TOC \\o "1-3" \\h \\z \\u',
                        'Mục lục sẽ hiện ở đây sau khi cập nhật trường (F9).')
        set_run_font(run, size=12)
        note = self.doc.add_paragraph()
        note.paragraph_format.space_before = Pt(10)
        set_run_font(note.add_run('Ghi chú khi hoàn thiện: bấm chuột phải vào vùng mục lục → '
                                  'Update Field (hoặc F9) để Word tính lại số trang sau khi '
                                  'đã chèn toàn bộ hình và bảng.'),
                     size=10.5, italic=True, color=RGBColor(0x77, 0x77, 0x77))
        self.reset_num()


# ---------------------------------------------------------------- cover
def build_cover(doc, cover):
    para(doc, cover['org1'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=18, space_after=2)
    para(doc, cover['org2'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=22)
    para(doc, cover['org3'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = line._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F3864')
    pbdr.append(bottom)
    pPr.append(pbdr)
    line.paragraph_format.space_after = Pt(24)

    para(doc, cover['kind'], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=26)
    para(doc, cover['label'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    para(doc, cover['title'], size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=26, line=1.35)
    para(doc, cover['field'], size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=10, space_after=34)
    para(doc, 'Tác giả:', size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    for name, role in cover['authors']:
        para(doc, name, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        para(doc, '(%s)' % role, size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_after=8)
    for _ in range(3):
        para(doc, '', space_after=10)
    para(doc, cover['place'], size=13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=36)


# ---------------------------------------------------------------- main
def load_blocks():
    mods, blocks = [], []
    for name in MODULES:
        try:
            m = importlib.import_module(name)
        except ModuleNotFoundError:
            continue
        mods.append(m)
        blocks += m.BLOCKS
    return mods, blocks


def number_blocks(blocks):
    fig_no, tab_no = {}, {}
    fig_titles, tab_titles = {}, {}
    for b in blocks:
        if b[0] == 'figure' and b[1] not in fig_no:
            fig_no[b[1]] = len(fig_no) + 1
            fig_titles[fig_no[b[1]]] = b[2]
        if b[0] == 'table' and b[1] not in tab_no:
            tab_no[b[1]] = len(tab_no) + 1
            tab_titles[tab_no[b[1]]] = b[2]
    return fig_no, tab_no, fig_titles, tab_titles


def main():
    mods, blocks = load_blocks()
    fig_no, tab_no, fig_titles, tab_titles = number_blocks(blocks)

    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    build_cover(doc, mods[0].COVER)

    b = Builder(doc, fig_no, tab_no, fig_titles, tab_titles)
    for block in blocks:
        kind = block[0]
        if kind == 'h1':
            b.h1(block[1])
        elif kind == 'h1c':
            b.h1c(block[1])
        elif kind == 'h2':
            b.h2(block[1])
        elif kind == 'h3':
            b.h3(block[1])
        elif kind == 'p':
            b.p(block[1])
        elif kind == 'label':
            b.label(block[1])
        elif kind == 'mt':
            b.mt(block[1], block[2])
        elif kind == 'bullet':
            b.bullet(block[1])
        elif kind == 'numbered':
            b.numbered(block[1])
        elif kind == 'note':
            b.note(block[1])
        elif kind == 'spacer':
            b.spacer()
        elif kind == 'right':
            b.right(block[1], block[2])
        elif kind == 'table':
            b.table(block[1], block[2], block[3], block[4], block[5])
        elif kind == 'figure':
            b.figure(block[1], block[2], block[3])
        elif kind == 'code':
            b.code(block[1], block[2])
        elif kind == 'pre':
            b.pre(block[1])
        elif kind == 'refs':
            b.refs(block[1])
        elif kind == 'toc':
            b.toc(block[1])
        elif kind == 'figlist':
            b.figlist(block[1])
        elif kind == 'tablist':
            b.tablist(block[1])
        elif kind == 'pagebreak':
            doc.add_page_break()
        else:
            raise ValueError('Block không hợp lệ: %r' % (kind,))

    out = os.environ.get('OUT', os.path.join(HERE, '..', 'báo cáo', 'BAO_CAO_ELEC_AI.docx'))
    out = os.path.abspath(out)
    doc.save(out)
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    words += sum(len(c.text.split()) for t in doc.tables for r in t.rows for c in r.cells)
    print('saved:', out)
    print('blocks:', len(blocks), '| figures:', len(fig_no), '| tables:', len(tab_no),
          '| words:', words)


if __name__ == '__main__':
    main()

