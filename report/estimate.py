# -*- coding: utf-8 -*-
"""Ước lượng số trang của báo cáo từ cấu trúc tệp docx."""
import math
import sys
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

PATH = sys.argv[1] if len(sys.argv) > 1 else r'D:\Tải Xuống\app\báo cáo\BAO_CAO_ELEC_AI.docx'
TEXT_H = 24.7  # cm vùng chữ
TEXT_W = 16.0

d = docx.Document(PATH)


def iter_block_items(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


def para_height(p):
    text = p.text or ''
    if not text.strip():
        return 0.35
    size = 13
    for r in p.runs:
        if r.font.size:
            size = r.font.size.pt
            break
    if p.style.name.startswith('Heading'):
        size = {'Heading 1': 16, 'Heading 2': 14, 'Heading 3': 13}.get(p.style.name, 13)
    chars_per_line = TEXT_W / (size * 0.0205)  # ~ cm trên ký tự
    line_h = size * 0.0353 * 1.3  # cm
    lines = max(1, math.ceil(len(text) / max(1, chars_per_line)))
    extra = 0.0
    if p.paragraph_format.space_before:
        extra += p.paragraph_format.space_before.pt * 0.0353
    if p.paragraph_format.space_after:
        extra += p.paragraph_format.space_after.pt * 0.0353
    if p.style.name == 'Heading 1':
        extra += 0.3
    return lines * line_h + extra


def table_height(t):
    total = 0.0
    ncols = len(t.columns)
    widths = []
    for c in t.rows[0].cells:
        widths.append(c.width.cm if c.width else TEXT_W / ncols)
    if not widths or sum(widths) < 1:
        widths = [TEXT_W / ncols] * ncols
    scale = TEXT_W / sum(widths)
    widths = [w * scale for w in widths]
    for row in t.rows:
        h = 0.30
        for i, cell in enumerate(row.cells):
            txt = cell.text or ''
            if not txt.strip():
                continue
            # font nhỏ hơn với bảng nhiều cột
            fs = 12 if ncols <= 5 else (10.5 if ncols == 6 else 9.5)
            cpl = max(4, widths[i] / (fs * 0.0205))
            lines = 0
            for part in txt.split('\n'):
                lines += max(1, math.ceil(len(part) / cpl))
            h = max(h, lines * fs * 0.0353 * 1.05 + 0.14)
        total += h
    return total + 0.3


total_h = 0.0
fig_boxes = 0
for item in iter_block_items(d):
    if isinstance(item, Paragraph):
        total_h += para_height(item)
        if 'CHÈN HÌNH' in (item.text or ''):
            fig_boxes += 1
    else:
        # bảng 1 ô lớn = khung hình
        if len(item.rows) == 1 and len(item.columns) == 1 and 'CHÈN HÌNH' in item.cell(0, 0).text:
            total_h += 8.9
        else:
            total_h += table_height(item)

pages = total_h / TEXT_H
print('chiều cao ước lượng (cm):', round(total_h, 1))
print('số trang ước lượng:', round(pages, 1))
print('khung hình:', fig_boxes, '| bảng dữ liệu:', len(d.tables) - fig_boxes)
