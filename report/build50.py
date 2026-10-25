# -*- coding: utf-8 -*-
"""Bản gọn ~50 trang: lọc bớt mục, đánh số lại hình/bảng, giữ nguyên khung báo cáo."""

import importlib
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build as B  # noqa: E402


DROP_CODE = ['vòng khóa dự phòng', 'các ngưỡng quyết định có gọi AI',
             'nguyên lí nhận dạng', 'sinh mã lớp ngẫu nhiên']
DROP_FIG_CAPTIONS = [
    'Cấu trúc thư mục mã nguồn của dự án',
    'Trang quiz và trang sổ câu sai',
    'Biểu đồ kết quả khảo sát người dùng',
    'Quy trình đóng gói và triển khai hệ thống',
    'Giao diện làm bài thi trực tuyến của học sinh',
]
DROP_TABLE_CAPTIONS = [
    'Bảng 6. Các nhóm biến cấu hình của hệ thống (giá trị minh họa)',
]
SKIP_SECTIONS = ['4.11.', '3.5.', '5.5.', '7.6.']


FRONT = [
    ('pagebreak',),
    ('h1c', 'LỜI CẢM ƠN'),
    ('p', 'Nhóm chúng em xin cảm ơn Ban Giám hiệu nhà trường đã tạo điều kiện về thời gian và không gian để nhóm hoàn thành dự án. Nhóm cảm ơn các thầy cô trong tổ bộ môn Hóa học đã đọc bản thảo và góp ý về cách trình bày lời giải, cách chia mức độ trong ma trận đề và những chỗ chưa đúng với chương trình hiện hành.'),
    ('p', 'Nhóm cũng cảm ơn các bạn trong câu lạc bộ tin học của trường đã dùng thử sản phẩm ở giai đoạn đầu và ghi lại những chỗ khó hiểu. Nhiều lỗi được nêu trong báo cáo này xuất phát từ những buổi thử đó.'),
    ('p', 'Dù đã cố gắng, sản phẩm và báo cáo vẫn còn những hạn chế đã ghi rõ ở Phần III. Nhóm mong nhận được góp ý của hội đồng để tiếp tục hoàn thiện.'),
    ('spacer',),
    ('right', 'Nhóm tác giả', True),
    ('right', 'Nguyễn Thế Duy – Nguyễn Bùi Nguyệt Minh', False),
    ('pagebreak',),
    ('h1c', 'CAM KẾT CỦA NHÓM TÁC GIẢ'),
    ('p', 'Nhóm chúng em cam kết những điều sau về báo cáo và sản phẩm đi kèm:'),
    ('numbered', 'Sản phẩm mô tả trong báo cáo là sản phẩm thật, chạy được tại địa chỉ đã nêu, không phải bản mô phỏng giao diện.'),
    ('numbered', 'Toàn bộ mã nguồn do hai thành viên trong nhóm viết, có lịch sử phiên bản làm bằng chứng.'),
    ('numbered', 'Các số liệu về khối lượng mã nguồn, số lần ghi nhận thay đổi và kết quả kiểm thử được lấy trực tiếp từ kho mã nguồn và từ lần chạy ngày 24 tháng 10 năm 2026.'),
    ('numbered', 'Những nội dung chưa có số liệu thật được đánh dấu rõ bằng cụm từ in nghiêng trong ngoặc vuông, không được viết thành số liệu đã đo.'),
    ('numbered', 'Báo cáo không chứa khóa truy cập, mật khẩu hay thông tin tài khoản của bất kì cá nhân, tổ chức nào; các đoạn mã trích dẫn đã thay bằng giá trị mẫu.'),
    ('spacer',),
    ('right', 'Hưng Yên, ngày ...... tháng ...... năm 2026', False),
    ('right', 'Đại diện nhóm tác giả', True),
    ('right', '(Kí và ghi rõ họ tên)', False),
    ('spacer',),
    ('spacer',),
    ('right', 'Nguyễn Thế Duy', False),
]


class ShortBuilder(B.Builder):
    def right(self, text, bold=False):
        B.para(self.doc, text, bold=bold, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4)
        self.reset_num()

    def spacer(self):
        B.para(self.doc, '', space_after=10)
        self.reset_num()


def preprocess(blocks):
    kept = []
    skip = False
    for b in blocks:
        k = b[0]
        if skip:
            if k in ('h2', 'h3'):
                skip = False
            else:
                continue
        if k == 'h3' and any(b[1].startswith(s) for s in SKIP_SECTIONS):
            skip = True
            continue
        if k == 'code' and any(x in b[1] for x in DROP_CODE):
            continue
        if k == 'figure' and b[2] in DROP_FIG_CAPTIONS:
            continue
        if k == 'table' and b[1] in DROP_TABLE_CAPTIONS:
            continue
        kept.append(b)

    fmap, tmap = {}, {}
    fn = tn = 0
    for b in kept:
        if b[0] == 'figure':
            fn += 1
            fmap[int(re.match(r'Hình (\d+)', b[1]).group(1))] = fn
        elif b[0] == 'table':
            tn += 1
            tmap[int(re.match(r'Bảng (\d+)', b[1]).group(1))] = tn

    def repl(s):
        s = re.sub(r'([Hh])ình (\d+)',
                   lambda m: '%sình %d' % (m.group(1), fmap.get(int(m.group(2)), int(m.group(2)))), s)
        s = re.sub(r'([Bb])ảng (\d+)',
                   lambda m: '%sảng %d' % (m.group(1), tmap.get(int(m.group(2)), int(m.group(2)))), s)
        return s

    out = []
    for b in kept:
        if b[0] == 'figure':
            n = fmap[int(re.match(r'Hình (\d+)', b[1]).group(1))]
            out.append(('figure', 'Hình %d' % n, repl(b[2]), b[3]))
        elif b[0] == 'table':
            m = re.match(r'Bảng (\d+)\.(.*)', b[1])
            cap = 'Bảng %d.%s' % (tmap[int(m.group(1))], m.group(2))
            hdr = [repl(h) for h in b[2]]
            rows = [[repl(str(c)) for c in row] for row in b[3]]
            out.append(('table', cap, hdr, rows, b[4]))
        else:
            lst = [repl(x) if isinstance(x, str) else x for x in b]
            out.append(tuple(lst))
    return out


def _render(b, block):
    kind = block[0]
    if kind == 'h1':
        b.h1(block[1])
    elif kind == 'h1c':
        b.h1c(block[1])
    elif kind == 'h2':
        b.h2(block[1])
    elif kind == 'h3':
        b.h3(block[1])
    elif kind == 'p':
        b.p(block[1])
    elif kind == 'label':
        b.label(block[1])
    elif kind == 'mt':
        b.mt(block[1], block[2])
    elif kind == 'bullet':
        b.bullet(block[1])
    elif kind == 'numbered':
        b.numbered(block[1])
    elif kind == 'table':
        b.table(block[1], block[2], block[3], block[4])
    elif kind == 'figure':
        b.figure(block[1], block[2], block[3])
    elif kind == 'code':
        b.code(block[1], block[2])
    elif kind == 'pre':
        b.pre(block[1])
    elif kind == 'refs':
        b.refs(block[1])
    elif kind == 'toc':
        b.toc(block[1])
    elif kind == 'pagebreak':
        b.doc.add_page_break()
    elif kind == 'right':
        b.right(block[1], block[2])
    elif kind == 'spacer':
        b.spacer()
    else:
        raise ValueError(kind)


def main():
    mods = {i: importlib.import_module('content%d' % i) for i in (1, 2, 4, 5)}
    c3a = importlib.import_module('content3a')
    c3b = importlib.import_module('content3b')

    blocks = []
    blocks += mods[1].BLOCKS
    blocks += mods[2].BLOCKS
    blocks += c3a.BLOCKS
    blocks += c3b.BLOCKS
    blocks += mods[4].BLOCKS
    blocks += mods[5].BLOCKS
    blocks = preprocess(blocks)

    doc = Document()
    B.setup_page(doc)
    B.setup_styles(doc)
    B.build_cover(doc, mods[1].COVER)

    b = ShortBuilder(doc)
    inserted = False
    for block in blocks:
        if not inserted and block[0] == 'h1' and block[1].startswith('PHẦN I'):
            for fb in FRONT:
                _render(b, fb)
            inserted = True
        _render(b, block)

    out = os.path.abspath(os.path.join(HERE, '..', 'báo cáo', 'BAO_CAO_ELEC_AI_50TRANG.docx'))
    doc.save(out)

    words = sum(len(p.text.split()) for p in doc.paragraphs)
    words += sum(len(c.text.split()) for t in doc.tables for r in t.rows for c in r.cells)
    print('saved:', out)
    print('blocks:', len(blocks), '| words:', words, '| tables:', len(doc.tables))


if __name__ == '__main__':
    main()
